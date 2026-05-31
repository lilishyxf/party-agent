# Convert test report markdown to docx (V4 style).
# Reads Chinese from external .md file to avoid encoding issues.
import re
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

INPUT = "D:/Projects/party-agent 2/outputs/测试报告_完整内容.md"
OUTPUT = "D:/Projects/party-agent 2/outputs/党务Agent_系统测试报告_V1.docx"

SAN_HAO = Pt(16)
SI_HAO = Pt(14)
XIAO_SI = Pt(12)
WU_HAO = Pt(10.5)

with open(INPUT, "r", encoding="utf-8") as f:
    raw = f.read()

lines = raw.split("\n")
blocks = []
i = 0
while i < len(lines):
    s = lines[i].strip()
    if s == "" or s == "---":
        i += 1; continue
    if s.startswith("# ") and not s.startswith("## "):
        blocks.append(("h1", s[2:].strip()))
    elif s.startswith("## "):
        blocks.append(("h2", s[3:].strip()))
    elif s.startswith("### "):
        blocks.append(("h3", s[4:].strip()))
    elif s.startswith("#### "):
        blocks.append(("h3", s[5:].strip()))
    elif s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        data = [r for r in rows if not all(re.match(r"^-+$", c) for c in r)]
        blocks.append(("table", data))
        continue
    else:
        blocks.append(("p", s))
    i += 1

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.17)
sec.right_margin = Cm(3.17)

# Footer: page number
ftr = sec.footer
ftr.is_linked_to_previous = False
fp = ftr.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run()._r.append(parse_xml('<w:fldChar %s w:fldCharType="begin"/>' % nsdecls("w")))
r2 = fp.add_run()
r2._r.append(parse_xml('<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls("w")))
fp.add_run()._r.append(parse_xml('<w:fldChar %s w:fldCharType="end"/>' % nsdecls("w")))

# Header
hdr = sec.header
hdr.is_linked_to_previous = False
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run("西安工业大学毕业设计(论文)")
hr.font.size = WU_HAO
hr.font.name = "Times New Roman"
hr._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
pPr = hp._p.get_or_add_pPr()
pPr.append(parse_xml('<w:pBdr %s><w:bottom w:val="double" w:sz="6" w:space="1" w:color="auto"/></w:pBdr>' % nsdecls("w")))

def sf(run, size=XIAO_SI, bold=False):
    run.font.size = size
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.bold = bold

def spf(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None, sb=Pt(0), sa=Pt(0)):
    pf = para.paragraph_format
    pf.alignment = align
    pf.line_spacing = Pt(18)
    pf.line_spacing_rule = 0
    pf.space_before = sb
    pf.space_after = sa
    if indent:
        pf.first_line_indent = indent

def h1(text):
    p = doc.add_paragraph()
    spf(p, WD_ALIGN_PARAGRAPH.CENTER, sb=Pt(24), sa=Pt(24))
    r = p.add_run(text)
    sf(r, SAN_HAO, True)

def h2(text):
    p = doc.add_paragraph()
    spf(p, WD_ALIGN_PARAGRAPH.LEFT, sb=Pt(12), sa=Pt(6))
    r = p.add_run(text)
    sf(r, SI_HAO, True)

def h3(text):
    p = doc.add_paragraph()
    spf(p, WD_ALIGN_PARAGRAPH.LEFT, sb=Pt(6), sa=Pt(3))
    r = p.add_run(text)
    sf(r, XIAO_SI, True)

def para(text):
    p = doc.add_paragraph()
    spf(p, indent=Emu(480 * 2 * 635))
    r = p.add_run(text)
    sf(r)

def table(data):
    hdr_row = data[0]
    rows = data[1:]
    nc = len(hdr_row)
    tbl = doc.add_table(rows=1 + len(rows), cols=nc)
    tbl.style = "Table Grid"
    for j, ct in enumerate(hdr_row):
        c = tbl.rows[0].cells[j]
        c.text = ""
        pp = c.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(ct)
        sf(rr, WU_HAO, True)
        shd = parse_xml('<w:shd %s w:fill="D9E2F3"/>' % nsdecls("w"))
        c._tc.get_or_add_tcPr().append(shd)
    for i, row in enumerate(rows):
        for j, ct in enumerate(row):
            if j < nc:
                c = tbl.rows[i + 1].cells[j]
                c.text = ""
                pp = c.paragraphs[0]
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = pp.add_run(ct)
                sf(rr, WU_HAO)

# Render
for btype, content in blocks:
    if btype == "h1":
        h1(content)
    elif btype == "h2":
        h2(content)
    elif btype == "h3":
        h3(content)
    elif btype == "p":
        para(content)
    elif btype == "table":
        table(content)

doc.save(OUTPUT)
print("Done: " + OUTPUT)
